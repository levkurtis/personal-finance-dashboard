"""Generate a fake "Financial <year>" workbook with the same layout as the real Google Sheet.

All values are random-but-plausible from a seeded RNG — never derived from real numbers.
"""

import random
from io import BytesIO

from openpyxl import Workbook

MONTHS = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]
# The real sheet has trailing spaces on some month names.
TRAILING_SPACE_MONTHS = {3, 10}

# Header row B..L, exactly as in the real sheet. Maps header -> source key (None = not a source).
HEADERS: list[tuple[str, str | None]] = [
    ("Total Income (Netto)", None),
    ("Pension", "pension"),
    ("Salary (Brutto)", "salary_brutto"),
    ("Salary (Netto)", "salary_netto"),
    ("Stock program", "stock_program"),
    ("Interest", "interest"),
    ("Tax", "tax"),
    ("Side income", "side_income"),
    ("Stocks Realization", "stocks_realization"),
    ("Bank benefits", "bank_benefits"),
    ("Income avg.", None),
]
NETTO_KEYS = [
    "salary_netto", "stock_program", "interest", "tax",
    "side_income", "stocks_realization", "bank_benefits",
]


def _generate(rng: random.Random) -> dict[int, dict[str, int]]:
    """Amounts in øre per month per source. 0 means the cell is left blank."""
    months: dict[int, dict[str, int]] = {m: {k: 0 for _, k in HEADERS if k} for m in range(1, 13)}

    brutto = rng.randrange(50_000, 54_000) * 100
    raise_month = rng.randint(5, 8)
    raise_by = rng.randrange(2_000, 4_000) * 100
    netto_share = rng.uniform(0.55, 0.62)
    pension_share = rng.uniform(0.08, 0.12)

    for m, row in months.items():
        if m == raise_month:
            brutto = min(brutto + raise_by, 5_800_000)
        row["salary_brutto"] = brutto
        row["salary_netto"] = round(brutto * (netto_share + rng.uniform(-0.01, 0.01)))
        row["pension"] = round(brutto * pension_share / 100) * 100
        if m <= 4:
            row["stock_program"] = rng.randrange(2_000, 4_500) * 100
        if rng.random() < 0.7:
            row["interest"] = rng.randrange(1, 40_000)

    tax_month = rng.randint(3, 6)
    months[tax_month]["tax"] = rng.choice([-1, 1]) * rng.randrange(500, 8_000) * 100
    for m in rng.sample(range(1, 13), 3):
        months[m]["side_income"] = rng.randrange(100, 2_500) * 100
    for m in rng.sample(range(1, 13), rng.randint(0, 2)):
        months[m]["stocks_realization"] = rng.randrange(5_000, 40_000) * 100 + rng.randrange(100)
    months[rng.randint(1, 12)]["bank_benefits"] = rng.randrange(500, 2_500) * 100
    return months


def build_demo_workbook(year: int = 2026, seed: int = 42) -> tuple[Workbook, dict]:
    """Return (workbook, expected). `expected` is the ground truth in øre:

    {"year": 2026, "months": {1: {"salary_netto": ..., ...}, ...}, "sheet_netto": {1: ..., ...}}
    """
    rng = random.Random(seed)
    months = _generate(rng)
    sheet_netto = {m: sum(row[k] for k in NETTO_KEYS) for m, row in months.items()}

    wb = Workbook()
    ws = wb.active
    ws.title = "Income"
    ws["A1"] = str(year)  # text, like the real sheet
    for col, (header, _) in enumerate(HEADERS, start=2):
        ws.cell(row=1, column=col, value=header)

    totals = [0] * len(HEADERS)
    for m in range(1, 13):
        r = m + 1
        name = MONTHS[m - 1] + (" " if m in TRAILING_SPACE_MONTHS else "")
        ws.cell(row=r, column=1, value=name)
        for i, (header, key) in enumerate(HEADERS):
            col = i + 2
            if header == "Total Income (Netto)":
                ore = sheet_netto[m]  # a number, not a formula: openpyxl can't write cached values
            elif header == "Income avg.":
                ore = rng.randrange(1_000_000, 5_000_000) if rng.random() < 0.6 else 0
            else:
                ore = months[m][key]
            if ore:
                ws.cell(row=r, column=col, value=ore / 100)
                totals[i] += ore

    ws.cell(row=14, column=1, value="Total")
    for i, (header, _) in enumerate(HEADERS):
        if header != "Income avg.":
            ws.cell(row=14, column=i + 2, value=totals[i] / 100)

    junk = wb.create_sheet("Expenses")
    junk["A1"] = "Rent"
    junk["B1"] = rng.randrange(5_000, 12_000)
    junk["A2"] = "January"
    junk["B2"] = "not income"

    expected = {"year": year, "months": months, "sheet_netto": sheet_netto}
    return wb, expected


def workbook_to_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
