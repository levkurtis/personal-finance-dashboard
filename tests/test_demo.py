from io import BytesIO

from openpyxl import load_workbook

from finance.demo import HEADERS, build_demo_workbook, workbook_to_bytes


def _reload(wb):
    return load_workbook(BytesIO(workbook_to_bytes(wb)), data_only=True)


def test_same_seed_same_values():
    _, a = build_demo_workbook(seed=7)
    _, b = build_demo_workbook(seed=7)
    _, c = build_demo_workbook(seed=8)
    assert a == b
    assert a != c


def test_layout_matches_real_sheet():
    wb, expected = build_demo_workbook(2026)
    wb = _reload(wb)
    assert wb.sheetnames == ["Income", "Expenses"]
    ws = wb["Income"]
    assert ws["A1"].value == "2026"
    assert [ws.cell(1, c).value for c in range(2, 13)] == [h for h, _ in HEADERS]
    names = [ws.cell(r, 1).value for r in range(2, 14)]
    assert names[0] == "January" and names[11] == "December"
    assert any(n.endswith(" ") for n in names)
    assert ws["A14"].value == "Total"
    assert expected["year"] == 2026 and sorted(expected["months"]) == list(range(1, 13))


def test_values_are_plausible():
    _, expected = build_demo_workbook()
    months = expected["months"]
    for m, row in months.items():
        assert 5_000_000 <= row["salary_brutto"] <= 5_800_000
        assert 0 < row["salary_netto"] < row["salary_brutto"]
        assert 0 < row["pension"] < row["salary_brutto"]
        assert (row["stock_program"] > 0) == (m <= 4)
    assert sum(1 for r in months.values() if r["tax"]) == 1
    assert sum(1 for r in months.values() if r["bank_benefits"]) == 1


def test_blank_cells_and_sheet_netto():
    wb, expected = build_demo_workbook()
    ws = _reload(wb)["Income"]
    stock_program_col = 2 + [h for h, _ in HEADERS].index("Stock program")
    assert ws.cell(row=13, column=stock_program_col).value is None  # December stock program is blank, not 0
    for m in range(1, 13):
        assert round(ws.cell(row=m + 1, column=2).value * 100) == expected["sheet_netto"][m]
