"""Build small Income workbooks in memory for tests. Never use real data here."""

from openpyxl import Workbook

from finance.demo import workbook_to_bytes

STANDARD_HEADERS = [
    "Total Income (Netto)", "Pension", "Salary (Brutto)", "Salary (Netto)", "Stock program",
    "Interest", "Tax", "Side income", "Stocks Realization",
    "Bank benefits", "Income avg.",
]


def make_income_xlsx(year_cell, headers, rows, sheet_name="Income", extra_sheets=()) -> bytes:
    """headers: row 1 from column B. rows: lists of [label, value_B, value_C, ...]."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws["A1"] = year_cell
    for col, header in enumerate(headers, start=2):
        ws.cell(row=1, column=col, value=header)
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            if value is not None:
                ws.cell(row=r, column=c, value=value)
    for name in extra_sheets:
        wb.create_sheet(name)["A1"] = "junk"
    return workbook_to_bytes(wb)
